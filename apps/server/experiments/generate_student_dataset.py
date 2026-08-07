from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import random
import re
import shutil
import string
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable, Optional

# 선택적 외부 라이브러리
try:
    from docx import Document
except ImportError:
    Document = None

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
except ImportError:
    A4 = None
    pdfmetrics = None
    TTFont = None
    canvas = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


# =========================================================
# 기본 설정
# =========================================================

DEFAULT_FILE_COUNT = 1000
DEFAULT_SEED = 20260722
DEFAULT_OUTPUT_ROOT = Path.home() / "student_dataset"

TOP_LEVEL_CATEGORIES = [
    "project",
    "assignment",
    "essay",
    "research",
    "lecture",
    "reference",
    "practice",
    "team_project",
]

# 문서 제목 끝에 붙는 문서 유형.
DOCUMENT_TYPES = [
    "보고서",
    "과제",
    "프로젝트",
    "실험 결과",
    "설계 문서",
    "조사 자료",
    "발표 초안",
    "논문 요약",
]

# 수행 형태. 문서 본문에 "수행 형태:" 줄로 명시된다.
WORK_MODES = ["개인", "팀"]

# (문서 유형, 수행 형태) -> category 결정 규칙.
#
# 2026-07-26 변경: 이전에는 category를 rng.choice로 무작위 배정했기 때문에
# 문서 내용과 통계적으로 독립이었고, 분류 정확도의 이론적 상한이 무작위 수준(12.5%)이었다.
# 이제 category는 문서에 명시된 두 신호로 완전히 결정된다.
#
#   - 어느 한 신호만으로는 부족하다. 문서 유형만 보고 최빈값을 찍으면 상한 68.75%.
#   - 두 신호를 모두 읽으면 상한 100%.
#   - 무작위 기준선은 12.5%.
#
# 표를 8x2 = 16칸으로 두고 각 category에 정확히 2칸씩 배정해 분포를 균등하게 유지한다.
CATEGORY_RULES: dict[tuple[str, str], str] = {
    ("과제", "개인"): "assignment",
    ("과제", "팀"): "team_project",
    ("보고서", "개인"): "essay",
    ("보고서", "팀"): "assignment",
    ("프로젝트", "개인"): "project",
    ("프로젝트", "팀"): "team_project",
    ("설계 문서", "개인"): "project",
    ("설계 문서", "팀"): "reference",
    ("실험 결과", "개인"): "practice",
    ("실험 결과", "팀"): "practice",
    ("조사 자료", "개인"): "research",
    ("조사 자료", "팀"): "research",
    ("발표 초안", "개인"): "lecture",
    ("발표 초안", "팀"): "lecture",
    ("논문 요약", "개인"): "reference",
    ("논문 요약", "팀"): "essay",
}


def resolve_category(document_type: str, work_mode: str) -> str:
    """문서 유형과 수행 형태로부터 category를 결정합니다."""
    try:
        return CATEGORY_RULES[(document_type, work_mode)]
    except KeyError:
        raise ValueError(
            f"CATEGORY_RULES에 정의되지 않은 조합입니다: ({document_type!r}, {work_mode!r})"
        ) from None

SUBJECTS = {
    "artificial_intelligence": [
        "machine_learning",
        "deep_learning",
        "natural_language_processing",
        "computer_vision",
        "reinforcement_learning",
    ],
    "computer_science": [
        "data_structure",
        "algorithm",
        "operating_system",
        "computer_architecture",
        "compiler",
    ],
    "database": [
        "relational_database",
        "sql",
        "nosql",
        "data_modeling",
        "query_optimization",
    ],
    "network": [
        "tcp_ip",
        "routing",
        "socket_programming",
        "wireless_network",
        "network_security",
    ],
    "software_engineering": [
        "requirements",
        "software_design",
        "testing",
        "devops",
        "version_control",
    ],
    "cyber_security": [
        "cryptography",
        "web_security",
        "malware_analysis",
        "digital_forensics",
        "access_control",
    ],
    "data_science": [
        "data_analysis",
        "visualization",
        "statistics",
        "data_mining",
        "big_data",
    ],
    "web_programming": [
        "frontend",
        "backend",
        "rest_api",
        "web_framework",
        "cloud_deployment",
    ],
    "mobile_programming": [
        "android",
        "ios",
        "cross_platform",
        "mobile_ui",
        "mobile_database",
    ],
    "embedded_system": [
        "iot",
        "sensor_network",
        "microcontroller",
        "real_time_system",
        "robotics",
    ],
}

FILE_TYPE_WEIGHTS = {
    ".txt": 16,
    ".md": 15,
    ".py": 13,
    ".pdf": 12,
    ".docx": 12,
    ".csv": 9,
    ".json": 7,
    ".html": 6,
    ".ipynb": 6,
    ".sql": 4,
}

KOREAN_TITLES = [
    "기계학습 모델의 성능 비교",
    "합성곱 신경망 기반 이미지 분류",
    "운영체제 프로세스 스케줄링 분석",
    "데이터베이스 정규화와 질의 최적화",
    "TCP 혼잡 제어 알고리즘 분석",
    "소프트웨어 요구사항 명세서",
    "웹 애플리케이션 보안 취약점 분석",
    "강화학습 기반 경로 탐색",
    "자연어 처리 모델의 문장 분류",
    "그래프 알고리즘 구현 및 실험",
    "클라우드 환경의 분산 처리 구조",
    "모바일 애플리케이션 설계",
    "사물인터넷 센서 데이터 분석",
    "컴퓨터 비전 객체 검출 실험",
    "암호화 알고리즘의 처리 성능 평가",
    "관계형 데이터베이스 설계 과제",
    "자료구조별 탐색 성능 비교",
    "REST API 서버 구현 보고서",
    "딥러닝 학습률 변화 실험",
    "빅데이터 처리 파이프라인 설계",
]

ENGLISH_KEYWORDS = [
    "machine learning",
    "deep learning",
    "computer vision",
    "natural language processing",
    "database",
    "algorithm",
    "network",
    "operating system",
    "software engineering",
    "cyber security",
    "cloud computing",
    "data science",
    "embedded system",
    "web programming",
    "mobile application",
]

STUDENT_NAMES = [
    "김민준",
    "김서연",
    "이도윤",
    "이하은",
    "박지호",
    "박수빈",
    "최현우",
    "최유진",
    "정우진",
    "정다은",
]

COURSE_NAMES = [
    "인공지능",
    "자료구조",
    "알고리즘",
    "운영체제",
    "컴퓨터네트워크",
    "데이터베이스",
    "소프트웨어공학",
    "정보보안",
    "웹프로그래밍",
    "컴퓨터구조",
]

SEMESTERS = [
    "2024_1학기",
    "2024_2학기",
    "2025_1학기",
    "2025_2학기",
    "2026_1학기",
]

DOCUMENT_SECTIONS = [
    "서론",
    "관련 연구",
    "문제 정의",
    "시스템 설계",
    "구현 방법",
    "실험 환경",
    "실험 결과",
    "분석 및 고찰",
    "결론",
    "참고문헌",
]

KOREAN_SENTENCES = [
    "본 과제에서는 주어진 문제를 해결하기 위한 알고리즘을 설계하고 성능을 분석하였다.",
    "실험 데이터는 가상으로 생성되었으며 실제 개인 정보나 연구 자료를 포함하지 않는다.",
    "모델의 학습 성능은 정확도, 정밀도, 재현율 및 F1 점수를 기준으로 평가하였다.",
    "데이터 전처리 과정에서는 결측값 처리, 정규화 및 학습 데이터 분할을 수행하였다.",
    "제안한 방법은 기준 모델과 비교하여 안정적인 성능을 나타냈다.",
    "프로그램은 Python을 기반으로 구현하였으며 재현 가능한 실험 환경을 구성하였다.",
    "실험 결과는 입력 데이터의 분포와 하이퍼파라미터 설정에 따라 달라질 수 있다.",
    "본 문서는 컴퓨터공학 전공 수업의 가상 과제 제출물을 모사하기 위해 작성되었다.",
    "시스템의 각 모듈은 독립적으로 동작하면서 정의된 인터페이스를 통해 데이터를 교환한다.",
    "향후 연구에서는 더 큰 데이터셋과 다양한 평가 환경을 적용할 필요가 있다.",
]

REFERENCE_TEMPLATES = [
    "Kim, J. and Lee, S. ({year}). A Study on {keyword}. Journal of Synthetic Computing, 12(3), 101-115.",
    "Park, H. ({year}). Experimental Analysis of {keyword}. Proceedings of the Virtual CS Conference, 44-51.",
    "Choi, M. and Jung, Y. ({year}). Design and Evaluation of {keyword}. Fake Academic Press.",
    "Lee, D. et al. ({year}). An Educational Approach to {keyword}. Computing Education Review, 8(2), 20-34.",
]


@dataclass
class FileMetadata:
    index: int
    file_name: str
    extension: str
    relative_path: str
    absolute_path: str
    category: str
    document_type: str
    work_mode: str
    subject: str
    subtopic: str
    semester: str
    course: str
    document_title: str
    size_bytes: int
    created_at: str
    modified_at: str
    sha256: str
    first_page_text: str
    extraction_status: str


# =========================================================
# 유틸리티
# =========================================================

def safe_filename(value: str, max_length: int = 90) -> str:
    """파일명에 사용할 수 없는 문자를 제거합니다."""
    value = re.sub(r'[\\/:*?"<>|]', "_", value)
    value = re.sub(r"\s+", "_", value.strip())
    return value[:max_length].strip("._") or "untitled"


def random_date(rng: random.Random) -> datetime:
    start = datetime(2024, 3, 1, 9, 0, 0)
    end = datetime(2026, 6, 30, 23, 59, 59)
    seconds = int((end - start).total_seconds())
    return start + timedelta(seconds=rng.randint(0, seconds))


def set_file_timestamp(path: Path, timestamp: datetime) -> None:
    epoch = timestamp.timestamp()
    path.touch(exist_ok=True)
    path.chmod(0o644)

    import os

    os.utime(path, (epoch, epoch))


def calculate_sha256(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def weighted_extension(rng: random.Random) -> str:
    extensions = list(FILE_TYPE_WEIGHTS.keys())
    weights = list(FILE_TYPE_WEIGHTS.values())
    return rng.choices(extensions, weights=weights, k=1)[0]


def find_korean_font() -> Optional[Path]:
    """운영체제에서 사용 가능한 한글 폰트를 검색합니다."""
    candidates = [
        Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJKkr-Regular.otf"),
        Path("/System/Library/Fonts/AppleSDGothicNeo.ttc"),
        Path("C:/Windows/Fonts/malgun.ttf"),
    ]

    for candidate in candidates:
        if candidate.exists():
            return candidate

    return None


def truncate_text(text: str, max_chars: int = 4000) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text[:max_chars]


# =========================================================
# 합성 문서 내용 생성
# =========================================================

def make_document_context(
    rng: random.Random,
    index: int,
    subject: str,
    subtopic: str,
    semester: str,
    course: str,
) -> dict:
    title = rng.choice(KOREAN_TITLES)
    keyword = rng.choice(ENGLISH_KEYWORDS)
    student_name = rng.choice(STUDENT_NAMES)
    student_id = f"20{rng.randint(20, 26):02d}{rng.randint(10000, 99999)}"

    title_suffix = rng.choice(DOCUMENT_TYPES)
    work_mode = rng.choice(WORK_MODES)

    # 팀 문서에는 팀원 명단을 함께 실어 "수행 형태" 신호를 본문에서 한 번 더 드러낸다.
    if work_mode == "팀":
        candidates = [name for name in STUDENT_NAMES if name != student_name]
        team_members = [student_name] + rng.sample(candidates, k=rng.randint(2, 3))
    else:
        team_members = [student_name]

    full_title = f"{title} {title_suffix}"

    return {
        "index": index,
        "title": full_title,
        "document_type": title_suffix,
        "work_mode": work_mode,
        "team_members": team_members,
        "subject": subject,
        "subtopic": subtopic,
        "semester": semester,
        "course": course,
        "student_name": student_name,
        "student_id": student_id,
        "keyword": keyword,
    }


def describe_work_mode(context: dict) -> str:
    """본문에 표시할 수행 형태 문구."""
    if context["work_mode"] == "팀":
        return f"팀 수행 ({len(context['team_members'])}인)"
    return "개인 수행"


def make_report_text(rng: random.Random, context: dict) -> str:
    lines = [
        context["title"],
        "",
        f"과목명: {context['course']}",
        f"학기: {context['semester']}",
        f"문서 유형: {context['document_type']}",
        f"수행 형태: {describe_work_mode(context)}",
        f"전공 분야: {context['subject']}",
        f"세부 주제: {context['subtopic']}",
        f"학번: {context['student_id']}",
        f"작성자: {context['student_name']}",
    ]

    if context["work_mode"] == "팀":
        lines.append(f"팀원: {', '.join(context['team_members'])}")

    lines += [
        "",
        "요약",
        (
            f"본 문서는 {context['keyword']} 주제에 관한 가상 과제 보고서이다. "
            "컴퓨터공학 전공 대학생의 파일 저장 구조를 모사하기 위한 합성 데이터로 작성되었다."
        ),
        "",
    ]

    selected_sections = rng.sample(
        DOCUMENT_SECTIONS,
        k=rng.randint(5, min(9, len(DOCUMENT_SECTIONS))),
    )

    for section_number, section in enumerate(selected_sections, start=1):
        lines.append(f"{section_number}. {section}")

        paragraph_count = rng.randint(2, 5)

        for _ in range(paragraph_count):
            sentence_count = rng.randint(3, 7)
            sentences = rng.choices(KOREAN_SENTENCES, k=sentence_count)
            lines.append(" ".join(sentences))

        lines.append("")

    lines.append("참고문헌")

    for _ in range(rng.randint(3, 8)):
        template = rng.choice(REFERENCE_TEMPLATES)
        reference = template.format(
            year=rng.randint(2018, 2026),
            keyword=context["keyword"].title(),
        )
        lines.append(reference)

    return "\n".join(lines)


def make_python_code(rng: random.Random, context: dict) -> str:
    sample_size = rng.randint(50, 300)
    return f'''"""
파일: {context["title"]}
과목: {context["course"]}
문서 유형: {context["document_type"]}
수행 형태: {describe_work_mode(context)}
작성자: {context["student_name"]}
팀원: {", ".join(context["team_members"])}
설명: 합성된 컴퓨터공학 과제용 Python 코드
"""

from __future__ import annotations

import random
from statistics import mean


def generate_data(size: int = {sample_size}, seed: int = {rng.randint(1, 9999)}) -> list[float]:
    random.seed(seed)
    return [random.uniform(0.0, 100.0) for _ in range(size)]


def normalize(values: list[float]) -> list[float]:
    if not values:
        return []

    minimum = min(values)
    maximum = max(values)

    if maximum == minimum:
        return [0.0 for _ in values]

    return [(value - minimum) / (maximum - minimum) for value in values]


def main() -> None:
    data = generate_data()
    normalized = normalize(data)

    print("주제:", "{context['subject']}/{context['subtopic']}")
    print("데이터 개수:", len(data))
    print("평균:", round(mean(data), 4))
    print("정규화 결과 앞부분:", normalized[:10])


if __name__ == "__main__":
    main()
'''


def make_sql_code(rng: random.Random, context: dict) -> str:
    return f"""-- {context['title']}
-- 과목: {context['course']}
-- 문서 유형: {context['document_type']}
-- 수행 형태: {describe_work_mode(context)}
-- 작성자: {context['student_name']}
-- 팀원: {", ".join(context['team_members'])}

CREATE TABLE students (
    student_id VARCHAR(20) PRIMARY KEY,
    student_name VARCHAR(50) NOT NULL,
    department VARCHAR(100) NOT NULL,
    grade INTEGER CHECK (grade BETWEEN 1 AND 4)
);

CREATE TABLE assignments (
    assignment_id INTEGER PRIMARY KEY,
    title VARCHAR(200) NOT NULL,
    subject VARCHAR(100),
    score DECIMAL(5, 2),
    submitted_at TIMESTAMP
);

INSERT INTO students
VALUES ('{context['student_id']}', '{context['student_name']}', '컴퓨터공학과', {rng.randint(1, 4)});

SELECT
    s.student_name,
    a.title,
    a.score
FROM students AS s
JOIN assignments AS a
    ON s.student_id = a.assignment_id::VARCHAR
WHERE a.subject = '{context['subject']}'
ORDER BY a.score DESC;
"""


def make_html_document(rng: random.Random, context: dict, report_text: str) -> str:
    escaped_text = html.escape(report_text)
    body = escaped_text.replace("\n", "<br>\n")

    return f"""<!DOCTYPE html>
<html lang="ko">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{html.escape(context['title'])}</title>
    <style>
        body {{
            max-width: 900px;
            margin: 40px auto;
            padding: 20px;
            font-family: sans-serif;
            line-height: 1.7;
        }}
        h1 {{
            border-bottom: 2px solid #333;
            padding-bottom: 12px;
        }}
    </style>
</head>
<body>
    <h1>{html.escape(context['title'])}</h1>
    <ul>
        <li>문서 유형: {html.escape(context['document_type'])}</li>
        <li>수행 형태: {html.escape(describe_work_mode(context))}</li>
        <li>팀원: {html.escape(", ".join(context['team_members']))}</li>
    </ul>
    <p>{body}</p>
</body>
</html>
"""


# =========================================================
# 파일 생성 함수
# =========================================================

def create_text_file(path: Path, report_text: str, _: dict, __: random.Random) -> None:
    path.write_text(report_text, encoding="utf-8")


def create_markdown_file(path: Path, report_text: str, context: dict, _: random.Random) -> None:
    lines = report_text.splitlines()
    markdown = [f"# {context['title']}", ""]

    for line in lines[1:]:
        if line in DOCUMENT_SECTIONS or re.match(r"^\d+\.\s", line):
            markdown.append(f"## {line}")
        else:
            markdown.append(line)

    path.write_text("\n".join(markdown), encoding="utf-8")


def create_python_file(path: Path, _: str, context: dict, rng: random.Random) -> None:
    path.write_text(make_python_code(rng, context), encoding="utf-8")


def create_sql_file(path: Path, _: str, context: dict, rng: random.Random) -> None:
    path.write_text(make_sql_code(rng, context), encoding="utf-8")


def create_json_file(path: Path, report_text: str, context: dict, rng: random.Random) -> None:
    payload = {
        "metadata": {
            "title": context["title"],
            "course": context["course"],
            "semester": context["semester"],
            "document_type": context["document_type"],
            "work_mode": describe_work_mode(context),
            "team_members": context["team_members"],
            "student_name": context["student_name"],
            "student_id": context["student_id"],
            "subject": context["subject"],
            "subtopic": context["subtopic"],
        },
        "experiment": {
            "accuracy": round(rng.uniform(0.65, 0.98), 4),
            "precision": round(rng.uniform(0.60, 0.97), 4),
            "recall": round(rng.uniform(0.58, 0.96), 4),
            "loss": round(rng.uniform(0.01, 0.75), 4),
            "epochs": rng.choice([10, 20, 30, 50, 100]),
        },
        "summary": truncate_text(report_text, 1200),
    }

    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def create_csv_file(path: Path, _: str, context: dict, rng: random.Random) -> None:
    # CSV에는 본문이 없으므로 문서 유형·수행 형태를 컬럼으로 실어 신호를 남긴다.
    # (기존 generator가 subject를 컬럼으로 넣던 방식과 동일한 패턴)
    fieldnames = [
        "sample_id",
        "feature_1",
        "feature_2",
        "feature_3",
        "label",
        "subject",
        "document_type",
        "work_mode",
    ]

    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()

        for row_index in range(rng.randint(30, 150)):
            writer.writerow(
                {
                    "sample_id": row_index + 1,
                    "feature_1": round(rng.uniform(-3, 3), 6),
                    "feature_2": round(rng.uniform(0, 100), 6),
                    "feature_3": rng.randint(0, 1000),
                    "label": rng.choice([0, 1, 2]),
                    "subject": context["subject"],
                    "document_type": context["document_type"],
                    "work_mode": describe_work_mode(context),
                }
            )


def create_html_file(path: Path, report_text: str, context: dict, rng: random.Random) -> None:
    path.write_text(
        make_html_document(rng, context, report_text),
        encoding="utf-8",
    )


def create_notebook_file(path: Path, report_text: str, context: dict, rng: random.Random) -> None:
    notebook = {
        "cells": [
            {
                "cell_type": "markdown",
                "metadata": {},
                "source": [
                    f"# {context['title']}\n",
                    f"- 과목: {context['course']}\n",
                    f"- 문서 유형: {context['document_type']}\n",
                    f"- 수행 형태: {describe_work_mode(context)}\n",
                    f"- 작성자: {context['student_name']}\n",
                    f"- 팀원: {', '.join(context['team_members'])}\n",
                    f"- 주제: {context['subject']}/{context['subtopic']}\n",
                ],
            },
            {
                "cell_type": "markdown",
                "metadata": {},
                "source": [truncate_text(report_text, 1800)],
            },
            {
                "cell_type": "code",
                "execution_count": None,
                "metadata": {},
                "outputs": [],
                "source": [
                    "import random\n",
                    f"random.seed({rng.randint(1, 9999)})\n",
                    "data = [random.random() for _ in range(100)]\n",
                    "sum(data) / len(data)\n",
                ],
            },
        ],
        "metadata": {
            "kernelspec": {
                "display_name": "Python 3",
                "language": "python",
                "name": "python3",
            },
            "language_info": {
                "name": "python",
                "version": "3.11",
            },
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }

    path.write_text(
        json.dumps(notebook, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def create_docx_file(path: Path, report_text: str, context: dict, _: random.Random) -> None:
    if Document is None:
        raise RuntimeError(
            "DOCX 생성을 위해 python-docx가 필요합니다. "
            "`pip install python-docx`를 실행하십시오."
        )

    document = Document()
    document.add_heading(context["title"], level=0)

    metadata_table = document.add_table(rows=0, cols=2)
    metadata = [
        ("과목명", context["course"]),
        ("학기", context["semester"]),
        ("문서 유형", context["document_type"]),
        ("수행 형태", describe_work_mode(context)),
        ("작성자", context["student_name"]),
        ("팀원", ", ".join(context["team_members"])),
        ("학번", context["student_id"]),
        ("전공 분야", context["subject"]),
        ("세부 주제", context["subtopic"]),
    ]

    for key, value in metadata:
        cells = metadata_table.add_row().cells
        cells[0].text = key
        cells[1].text = value

    document.add_paragraph("")

    for line in report_text.splitlines():
        stripped = line.strip()

        if not stripped:
            document.add_paragraph("")
        elif re.match(r"^\d+\.\s", stripped):
            document.add_heading(stripped, level=1)
        elif stripped in {"요약", "참고문헌"}:
            document.add_heading(stripped, level=1)
        else:
            document.add_paragraph(stripped)

    document.save(path)


def register_pdf_font() -> str:
    """PDF용 한글 폰트를 등록하고 폰트 이름을 반환합니다."""
    if pdfmetrics is None or TTFont is None:
        return "Helvetica"

    font_path = find_korean_font()

    if font_path is None:
        return "Helvetica"

    font_name = "SyntheticKoreanFont"

    try:
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
        return font_name
    except Exception:
        return "Helvetica"


def create_pdf_file(path: Path, report_text: str, context: dict, _: random.Random) -> None:
    if canvas is None or A4 is None:
        raise RuntimeError(
            "PDF 생성을 위해 reportlab이 필요합니다. "
            "`pip install reportlab`을 실행하십시오."
        )

    page_width, page_height = A4
    font_name = register_pdf_font()
    pdf = canvas.Canvas(str(path), pagesize=A4)

    left_margin = 55
    right_margin = 55
    top_margin = 55
    bottom_margin = 55
    line_height = 16
    usable_width = page_width - left_margin - right_margin

    def split_line(text: str, font_size: int = 10) -> list[str]:
        if not text:
            return [""]

        # 폰트에 따라 정확한 폭 계산이 다를 수 있어 보수적으로 분할
        max_chars = max(20, int(usable_width / (font_size * 0.65)))
        return [
            text[start : start + max_chars]
            for start in range(0, len(text), max_chars)
        ]

    y = page_height - top_margin

    pdf.setFont(font_name, 16)
    for line in split_line(context["title"], 16):
        pdf.drawString(left_margin, y, line)
        y -= 23

    pdf.setFont(font_name, 10)

    header_lines = [
        f"과목명: {context['course']}",
        f"학기: {context['semester']}",
        f"문서 유형: {context['document_type']}",
        f"수행 형태: {describe_work_mode(context)}",
        f"작성자: {context['student_name']}",
        f"팀원: {', '.join(context['team_members'])}",
        f"학번: {context['student_id']}",
        f"주제: {context['subject']} / {context['subtopic']}",
        "",
    ]

    all_lines = header_lines + report_text.splitlines()

    for source_line in all_lines:
        wrapped_lines = split_line(source_line, 10)

        for line in wrapped_lines:
            if y <= bottom_margin:
                pdf.showPage()
                pdf.setFont(font_name, 10)
                y = page_height - top_margin

            # Helvetica 사용 시 한글이 깨질 수 있어 ASCII 대체
            if font_name == "Helvetica":
                safe_line = line.encode("ascii", errors="replace").decode("ascii")
            else:
                safe_line = line

            pdf.drawString(left_margin, y, safe_line)
            y -= line_height

    pdf.save()


CREATORS: dict[str, Callable[[Path, str, dict, random.Random], None]] = {
    ".txt": create_text_file,
    ".md": create_markdown_file,
    ".py": create_python_file,
    ".sql": create_sql_file,
    ".json": create_json_file,
    ".csv": create_csv_file,
    ".html": create_html_file,
    ".ipynb": create_notebook_file,
    ".docx": create_docx_file,
    ".pdf": create_pdf_file,
}


# =========================================================
# 문서 텍스트 추출
# =========================================================

def extract_pdf_first_page(path: Path) -> tuple[str, str]:
    if PdfReader is None:
        return "", "pypdf_not_installed"

    try:
        reader = PdfReader(str(path))

        if not reader.pages:
            return "", "empty_pdf"

        text = reader.pages[0].extract_text() or ""
        return truncate_text(text), "success"

    except Exception as error:
        return "", f"pdf_error: {type(error).__name__}: {error}"


def extract_docx_first_page(path: Path) -> tuple[str, str]:
    """
    DOCX에는 PDF와 같은 고정 페이지 개념이 없습니다.
    따라서 첫 페이지에 해당한다고 간주할 문서 앞부분을 추출합니다.
    """
    if Document is None:
        return "", "python_docx_not_installed"

    try:
        document = Document(path)
        text_parts: list[str] = []

        for paragraph in document.paragraphs:
            value = paragraph.text.strip()

            if value:
                text_parts.append(value)

            if len("\n".join(text_parts)) >= 4000:
                break

        # 표 내용도 일부 포함
        if len("\n".join(text_parts)) < 4000:
            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )

                    if row_text:
                        text_parts.append(row_text)

                    if len("\n".join(text_parts)) >= 4000:
                        break

        return truncate_text("\n".join(text_parts)), "success"

    except Exception as error:
        return "", f"docx_error: {type(error).__name__}: {error}"


def extract_hwp_text(path: Path) -> tuple[str, str]:
    """
    기존 HWP 파일이 있을 때 `hwp5txt` 명령어로 텍스트를 추출합니다.

    Ubuntu/Debian 예:
        pip install pyhwp

    pyhwp 설치 후 hwp5txt 명령이 제공될 수 있습니다.
    """
    command = shutil.which("hwp5txt")

    if command is None:
        return "", "hwp5txt_not_installed"

    import subprocess

    try:
        result = subprocess.run(
            [command, str(path)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=30,
            check=False,
        )

        if result.returncode != 0:
            error_message = result.stderr.strip()
            return "", f"hwp5txt_error: {error_message}"

        return truncate_text(result.stdout), "success"

    except Exception as error:
        return "", f"hwp_error: {type(error).__name__}: {error}"


def extract_plain_text(path: Path) -> tuple[str, str]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
        return truncate_text(text), "success"
    except Exception as error:
        return "", f"text_error: {type(error).__name__}: {error}"


def extract_csv_preview(path: Path) -> tuple[str, str]:
    try:
        lines: list[str] = []

        with path.open("r", encoding="utf-8-sig", errors="replace") as file:
            for _ in range(30):
                line = file.readline()

                if not line:
                    break

                lines.append(line.rstrip("\n"))

        return truncate_text("\n".join(lines)), "success"

    except Exception as error:
        return "", f"csv_error: {type(error).__name__}: {error}"


def extract_notebook_preview(path: Path) -> tuple[str, str]:
    try:
        notebook = json.loads(path.read_text(encoding="utf-8"))
        text_parts: list[str] = []

        for cell in notebook.get("cells", []):
            source = cell.get("source", [])

            if isinstance(source, list):
                text_parts.append("".join(source))
            elif isinstance(source, str):
                text_parts.append(source)

            if len("\n".join(text_parts)) >= 4000:
                break

        return truncate_text("\n".join(text_parts)), "success"

    except Exception as error:
        return "", f"ipynb_error: {type(error).__name__}: {error}"


def extract_first_page_or_preview(path: Path) -> tuple[str, str]:
    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf_first_page(path)

    if extension == ".docx":
        return extract_docx_first_page(path)

    if extension == ".hwp":
        return extract_hwp_text(path)

    if extension == ".csv":
        return extract_csv_preview(path)

    if extension == ".ipynb":
        return extract_notebook_preview(path)

    if extension in {
        ".txt",
        ".md",
        ".py",
        ".sql",
        ".json",
        ".html",
        ".xml",
        ".yaml",
        ".yml",
    }:
        return extract_plain_text(path)

    return "", "unsupported_extension"


# =========================================================
# 데이터셋 생성
# =========================================================

def build_relative_directory(
    rng: random.Random,
    category: str,
    subject: str,
    subtopic: str,
    semester: str,
    course: str,
) -> Path:
    """
    폴더 깊이를 무작위로 달리하여 실제 학생 디렉터리처럼 구성합니다.
    """
    patterns = [
        Path(category) / semester / subject / subtopic,
        Path(category) / subject / course / subtopic,
        Path(semester) / category / subject / subtopic,
        Path(category) / subject / subtopic,
    ]

    return rng.choice(patterns)


def generate_unique_filename(
    index: int,
    context: dict,
    extension: str,
    used_paths: set[Path],
    directory: Path,
) -> str:
    title = safe_filename(context["title"])
    course = safe_filename(context["course"])

    candidates = [
        f"{index:04d}_{title}{extension}",
        f"{course}_{index:04d}_{title}{extension}",
        f"{context['student_id']}_{title}_{index:04d}{extension}",
    ]

    for candidate in candidates:
        candidate_path = directory / candidate

        if candidate_path not in used_paths:
            used_paths.add(candidate_path)
            return candidate

    random_token = "".join(random.choices(string.ascii_lowercase + string.digits, k=8))
    candidate = f"{index:04d}_{title}_{random_token}{extension}"
    used_paths.add(directory / candidate)
    return candidate


def create_dataset(
    output_root: Path,
    file_count: int,
    seed: int,
    overwrite: bool,
    legacy_random_category: bool = False,
) -> Path:
    rng = random.Random(seed)

    if output_root.exists() and overwrite:
        shutil.rmtree(output_root)

    output_root.mkdir(parents=True, exist_ok=True)

    metadata_rows: list[FileMetadata] = []
    used_paths: set[Path] = set()

    extensions = list(FILE_TYPE_WEIGHTS.keys())
    missing_dependencies: list[str] = []

    if Document is None and ".docx" in extensions:
        missing_dependencies.append("python-docx")

    if canvas is None and ".pdf" in extensions:
        missing_dependencies.append("reportlab")

    if missing_dependencies:
        raise RuntimeError(
            "다음 패키지가 설치되지 않았습니다: "
            + ", ".join(sorted(set(missing_dependencies)))
            + "\n설치 명령: pip install python-docx reportlab pypdf"
        )

    for index in range(1, file_count + 1):
        subject = rng.choice(list(SUBJECTS.keys()))
        subtopic = rng.choice(SUBJECTS[subject])
        semester = rng.choice(SEMESTERS)
        course = rng.choice(COURSE_NAMES)
        extension = weighted_extension(rng)

        context = make_document_context(
            rng=rng,
            index=index,
            subject=subject,
            subtopic=subtopic,
            semester=semester,
            course=course,
        )

        if legacy_random_category:
            # 구버전 재현용. 문서 내용과 무관한 무작위 라벨이라 평가에 쓸 수 없다.
            category = rng.choice(TOP_LEVEL_CATEGORIES)
        else:
            category = resolve_category(context["document_type"], context["work_mode"])

        report_text = make_report_text(rng, context)

        relative_directory = build_relative_directory(
            rng=rng,
            category=category,
            subject=subject,
            subtopic=subtopic,
            semester=semester,
            course=course,
        )

        target_directory = output_root / relative_directory
        target_directory.mkdir(parents=True, exist_ok=True)

        filename = generate_unique_filename(
            index=index,
            context=context,
            extension=extension,
            used_paths=used_paths,
            directory=target_directory,
        )

        file_path = target_directory / filename
        creator = CREATORS[extension]
        creator(file_path, report_text, context, rng)

        generated_time = random_date(rng)
        set_file_timestamp(file_path, generated_time)

        first_page_text, extraction_status = extract_first_page_or_preview(file_path)

        stat = file_path.stat()

        metadata_rows.append(
            FileMetadata(
                index=index,
                file_name=file_path.name,
                extension=file_path.suffix.lower().lstrip("."),
                relative_path=str(file_path.relative_to(output_root)),
                absolute_path=str(file_path.resolve()),
                category=category,
                document_type=context["document_type"],
                work_mode=context["work_mode"],
                subject=subject,
                subtopic=subtopic,
                semester=semester,
                course=course,
                document_title=context["title"],
                size_bytes=stat.st_size,
                created_at=datetime.fromtimestamp(stat.st_ctime).isoformat(
                    timespec="seconds"
                ),
                modified_at=datetime.fromtimestamp(stat.st_mtime).isoformat(
                    timespec="seconds"
                ),
                sha256=calculate_sha256(file_path),
                first_page_text=first_page_text,
                extraction_status=extraction_status,
            )
        )

        if index % 100 == 0 or index == file_count:
            print(f"[진행] {index}/{file_count}개 파일 생성 완료")

    metadata_path = output_root / "file_metadata.csv"
    write_metadata_csv(metadata_path, metadata_rows)

    summary_path = output_root / "dataset_summary.json"
    write_summary_json(
        summary_path=summary_path,
        output_root=output_root,
        metadata_rows=metadata_rows,
        seed=seed,
        legacy_random_category=legacy_random_category,
    )

    print()
    print("데이터셋 생성 완료")
    print(f"- 생성 루트: {output_root.resolve()}")
    print(f"- 일반 파일 수: {len(metadata_rows)}")
    print(f"- 메타데이터 CSV: {metadata_path.resolve()}")
    print(f"- 요약 JSON: {summary_path.resolve()}")

    return metadata_path


def write_metadata_csv(
    metadata_path: Path,
    metadata_rows: list[FileMetadata],
) -> None:
    fieldnames = list(FileMetadata.__annotations__.keys())

    # utf-8-sig를 사용하면 Excel에서 한글이 비교적 안정적으로 열립니다.
    with metadata_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()

        for row in metadata_rows:
            writer.writerow(row.__dict__)


def write_summary_json(
    summary_path: Path,
    output_root: Path,
    metadata_rows: list[FileMetadata],
    seed: int,
    legacy_random_category: bool = False,
) -> None:
    extension_counts: dict[str, int] = {}
    category_counts: dict[str, int] = {}
    subject_counts: dict[str, int] = {}
    extraction_counts: dict[str, int] = {}
    document_type_counts: dict[str, int] = {}
    work_mode_counts: dict[str, int] = {}

    for row in metadata_rows:
        extension_counts[row.extension] = extension_counts.get(row.extension, 0) + 1
        category_counts[row.category] = category_counts.get(row.category, 0) + 1
        subject_counts[row.subject] = subject_counts.get(row.subject, 0) + 1
        document_type_counts[row.document_type] = (
            document_type_counts.get(row.document_type, 0) + 1
        )
        work_mode_counts[row.work_mode] = work_mode_counts.get(row.work_mode, 0) + 1

        status_key = row.extraction_status.split(":", maxsplit=1)[0]
        extraction_counts[status_key] = extraction_counts.get(status_key, 0) + 1

    summary = {
        "dataset_root": str(output_root.resolve()),
        "generated_file_count": len(metadata_rows),
        "metadata_file": str((output_root / "file_metadata.csv").resolve()),
        "seed": seed,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "label_mode": "legacy_random" if legacy_random_category else "content_derived",
        "label_rule": (
            "category = rng.choice(TOP_LEVEL_CATEGORIES)  # 문서 내용과 독립"
            if legacy_random_category
            else "category = CATEGORY_RULES[(document_type, work_mode)]  # 문서 내용에서 결정"
        ),
        "extension_counts": extension_counts,
        "category_counts": category_counts,
        "subject_counts": subject_counts,
        "document_type_counts": document_type_counts,
        "work_mode_counts": work_mode_counts,
        "extraction_status_counts": extraction_counts,
        "notice": (
            "이 데이터셋은 실제 학생, 논문 또는 개인정보를 포함하지 않는 "
            "합성 데이터셋입니다."
        ),
    }

    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


# =========================================================
# 기존 폴더 스캔 기능
# =========================================================

def scan_existing_directory(
    target_root: Path,
    output_csv: Path,
) -> None:
    """
    이미 존재하는 폴더를 순회하면서 메타데이터 CSV를 생성합니다.
    PDF, DOCX, HWP 등은 앞부분 텍스트 추출을 시도합니다.
    """
    if not target_root.exists():
        raise FileNotFoundError(f"대상 폴더가 존재하지 않습니다: {target_root}")

    rows = []
    files = sorted(path for path in target_root.rglob("*") if path.is_file())

    for index, path in enumerate(files, start=1):
        first_page_text, extraction_status = extract_first_page_or_preview(path)
        stat = path.stat()

        rows.append(
            {
                "index": index,
                "file_name": path.name,
                "extension": path.suffix.lower().lstrip("."),
                "relative_path": str(path.relative_to(target_root)),
                "absolute_path": str(path.resolve()),
                "size_bytes": stat.st_size,
                "created_at": datetime.fromtimestamp(stat.st_ctime).isoformat(
                    timespec="seconds"
                ),
                "modified_at": datetime.fromtimestamp(stat.st_mtime).isoformat(
                    timespec="seconds"
                ),
                "sha256": calculate_sha256(path),
                "first_page_text": first_page_text,
                "extraction_status": extraction_status,
            }
        )

        if index % 100 == 0 or index == len(files):
            print(f"[스캔] {index}/{len(files)}개 파일 처리 완료")

    fieldnames = [
        "index",
        "file_name",
        "extension",
        "relative_path",
        "absolute_path",
        "size_bytes",
        "created_at",
        "modified_at",
        "sha256",
        "first_page_text",
        "extraction_status",
    ]

    output_csv.parent.mkdir(parents=True, exist_ok=True)

    with output_csv.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"스캔 메타데이터 저장 완료: {output_csv.resolve()}")


# =========================================================
# CLI
# =========================================================

def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "컴퓨터공학과 대학생의 과제·프로젝트 디렉터리를 모사한 "
            "합성 파일 데이터셋을 생성합니다."
        )
    )

    subparsers = parser.add_subparsers(dest="command")

    generate_parser = subparsers.add_parser(
        "generate",
        help="합성 데이터셋을 생성합니다.",
    )
    generate_parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_ROOT,
        help=f"출력 폴더. 기본값: {DEFAULT_OUTPUT_ROOT}",
    )
    generate_parser.add_argument(
        "--count",
        type=int,
        default=DEFAULT_FILE_COUNT,
        help=f"생성할 일반 파일 수. 기본값: {DEFAULT_FILE_COUNT}",
    )
    generate_parser.add_argument(
        "--seed",
        type=int,
        default=DEFAULT_SEED,
        help=f"난수 시드. 기본값: {DEFAULT_SEED}",
    )
    generate_parser.add_argument(
        "--overwrite",
        action="store_true",
        help="출력 폴더가 있으면 삭제하고 다시 생성합니다.",
    )
    generate_parser.add_argument(
        "--legacy-random-category",
        action="store_true",
        help=(
            "구버전(2026-07-26 이전) 동작 재현용. category를 문서 내용과 무관하게 "
            "무작위 배정합니다. 이 라벨은 모델 평가에 사용할 수 없습니다."
        ),
    )

    scan_parser = subparsers.add_parser(
        "scan",
        help="기존 디렉터리의 파일 정보를 CSV로 정리합니다.",
    )
    scan_parser.add_argument(
        "target",
        type=Path,
        help="스캔할 폴더",
    )
    scan_parser.add_argument(
        "--csv",
        type=Path,
        default=Path.cwd() / "scanned_file_metadata.csv",
        help="출력 CSV 경로",
    )

    return parser.parse_args()


def main() -> None:
    args = parse_arguments()

    if args.command in {None, "generate"}:
        output = getattr(args, "output", DEFAULT_OUTPUT_ROOT)
        count = getattr(args, "count", DEFAULT_FILE_COUNT)
        seed = getattr(args, "seed", DEFAULT_SEED)
        overwrite = getattr(args, "overwrite", False)
        legacy = getattr(args, "legacy_random_category", False)

        if count <= 0:
            raise ValueError("--count는 1 이상의 정수여야 합니다.")

        if legacy:
            print(
                "[경고] --legacy-random-category: category가 문서 내용과 무관하게 "
                "무작위 배정됩니다. 이 데이터셋의 category는 모델 평가에 쓸 수 없습니다."
            )

        create_dataset(
            output_root=output.expanduser(),
            file_count=count,
            seed=seed,
            overwrite=overwrite,
            legacy_random_category=legacy,
        )
        return

    if args.command == "scan":
        scan_existing_directory(
            target_root=args.target.expanduser(),
            output_csv=args.csv.expanduser(),
        )
        return

    raise ValueError(f"알 수 없는 명령입니다: {args.command}")


if __name__ == "__main__":
    main()